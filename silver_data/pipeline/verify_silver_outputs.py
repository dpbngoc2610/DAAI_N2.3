from __future__ import annotations

import argparse
import json
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


PARSER = argparse.ArgumentParser(description="Verify generated Silver workbooks.")
PARSER.add_argument("--manifest", type=Path, required=True)
PARSER.add_argument("--output-dir", type=Path, required=True)
PARSER.add_argument("--report", type=Path, required=True)
ARGS = PARSER.parse_args()

MANIFEST = json.loads(ARGS.manifest.resolve().read_text(encoding="utf-8"))
OUTPUT_DIR = ARGS.output_dir.resolve()
REPORT_PATH = ARGS.report.resolve()


results = []
all_ok = True

for meta in MANIFEST["datasets"]:
    name = meta["name"]
    path = OUTPUT_DIR / f"{name}_silver.xlsx"
    checks = {}
    checks["file_exists"] = path.exists() and path.stat().st_size > 0
    if not checks["file_exists"]:
        results.append({"name": name, "checks": checks})
        all_ok = False
        continue

    with zipfile.ZipFile(path) as zf:
        checks["zip_integrity"] = zf.testzip() is None

    wb = load_workbook(path, read_only=True, data_only=False)
    checks["sheet_names"] = wb.sheetnames == ["Silver", "Quality", "Quarantine"]
    ws = wb["Silver"]
    row_iterator = ws.iter_rows()
    header_cells = next(row_iterator)
    header = [cell.value for cell in header_cells]
    checks["headers"] = header == meta["silver_columns"]
    checks["quality_columns"] = "data_quality_status" in header and "warning_codes" in header
    status_idx = meta["silver_columns"].index("data_quality_status")
    code_idx = meta["silver_columns"].index("warning_codes")
    zip_idx = meta["silver_columns"].index("zip") if name in {"customers", "geography", "orders_enriched"} else None
    phone_idx = meta["silver_columns"].index("shipper_phone") if name == "shipments" else None
    first_data = []
    actual_rows = 1
    warning_rows = 0
    invalid_status_rows = 0
    sample_zips = []
    sample_phones = []
    for row_number, row in enumerate(row_iterator, start=2):
        actual_rows += 1
        if not first_data:
            first_data = list(row)
        status = row[status_idx].value
        codes = row[code_idx].value
        if status not in {"valid", "warning"}:
            invalid_status_rows += 1
        if status == "warning":
            warning_rows += 1
            if not codes:
                invalid_status_rows += 1
        elif codes:
            invalid_status_rows += 1
        if zip_idx is not None and row_number <= 202 and row[zip_idx].value is not None:
            sample_zips.append(row[zip_idx].value)
        if phone_idx is not None and row_number <= 101 and row[phone_idx].value is not None:
            sample_phones.append(row[phone_idx].value)

    checks["row_count"] = actual_rows == meta["silver_rows"] + 1
    checks["column_count"] = len(header) == len(meta["silver_columns"])
    type_checks = []
    for cell, col_name in zip(first_data, meta["silver_columns"]):
        dtype = meta["types"].get(col_name, "text")
        if cell.value is None:
            continue
        if dtype in {"date", "datetime"}:
            type_checks.append(isinstance(cell.value, datetime) and cell.data_type == "d")
        elif dtype in {"integer", "money", "decimal", "rate", "percent_points"}:
            type_checks.append(cell.data_type == "n")
        elif dtype == "boolean":
            type_checks.append(cell.data_type == "b")
        else:
            type_checks.append(cell.data_type in {"s", "inlineStr"})
    checks["representative_types"] = all(type_checks)

    q = wb["Quality"]
    checks["quality_counts"] = (
        q["B7"].value == meta["source_rows"]
        and q["B8"].value == meta["silver_rows"]
        and q["B9"].value == meta["quarantined_rows"]
        and q["B10"].value == meta.get("warning_rows", 0)
    )
    quarantine = wb["Quarantine"]
    quarantine_rows = quarantine.max_row if quarantine.max_row is not None else sum(1 for _ in quarantine.iter_rows())
    checks["quarantine_count"] = quarantine_rows == meta["quarantined_rows"] + 1

    if name in {"customers", "geography", "orders_enriched"}:
        checks["zip_as_five_char_text"] = all(isinstance(v, str) and len(v) == 5 and v.isdigit() for v in sample_zips)

    if name == "shipments":
        checks["phone_as_ten_char_text"] = all(isinstance(v, str) and len(v) == 10 and v.isdigit() for v in sample_phones)

    checks["warning_row_count"] = warning_rows == meta.get("warning_rows", 0)
    checks["warning_status_consistency"] = invalid_status_rows == 0

    wb.close()
    ok = all(checks.values())
    all_ok = all_ok and ok
    results.append({"name": name, "path": str(path), "bytes": path.stat().st_size, "checks": checks, "ok": ok})
    print(name, "PASS" if ok else "FAIL", checks, flush=True)

REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
REPORT_PATH.write_text(json.dumps({"all_ok": all_ok, "datasets": results}, ensure_ascii=False, indent=2), encoding="utf-8")
print(f"all_ok={all_ok}")
print(f"report={REPORT_PATH}")
raise SystemExit(0 if all_ok else 1)
